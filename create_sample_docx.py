from docx import Document

document = Document()

document.add_heading("Manjula Mallepeddi", level=1)

document.add_paragraph("Python Developer")

document.add_heading("Skills", level=2)
document.add_paragraph(
    "Python, FastAPI, SQL, Git, Generative AI, REST APIs"
)

document.add_heading("Experience", level=2)
document.add_paragraph(
    "Built an AI Resume Analyzer using FastAPI and Generative AI."
)

document.add_heading("Education", level=2)
document.add_paragraph(
    "Bachelor of Technology in Computer Science"
)

document.save("test/sample_resume.docx")

print("Sample DOCX created successfully.")